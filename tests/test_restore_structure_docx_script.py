# -*- coding: utf-8 -*-
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "src"))

from docx_builder import build_honest_docx
from restore_structure_docx import (
    build_best_docx,
    build_plain_docx_with_image_blocks,
    build_table_only_docx,
    extract_docx_pages,
    output_path,
)


def test_extract_docx_pages_accepts_page_lines_with_rect_fields(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "lines": [
                {"text": "left", "confidence": 0.99, "x1": 10, "y1": 20, "x2": 80, "y2": 40},
                {"text": "right", "confidence": 0.99, "x1": 900, "y1": 20, "x2": 980, "y2": 40},
            ],
        }],
    }

    pages = extract_docx_pages(payload)
    output = tmp_path / "restore.docx"
    output.write_bytes(build_honest_docx(pages, grid_cols=12, grid_rows=8))

    doc = Document(str(output))
    cells = [cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells]

    assert "left" in cells
    assert "right" in cells


def test_extract_docx_pages_can_use_overall_ocr_when_page_lines_are_missing():
    payload = {
        "pages": [{
            "page_no": 1,
            "structure": {
                "raw_results": [{
                    "res": {
                        "overall_ocr_res": {
                            "rec_texts": ["fallback"],
                            "rec_boxes": [[10, 20, 80, 40]],
                            "rec_scores": [0.98],
                        },
                    },
                }],
            },
        }],
    }

    pages = extract_docx_pages(payload)

    assert pages[0]["ocr_lines"][0]["text"] == "fallback"
    assert pages[0]["ocr_lines"][0]["bbox"] == [10.0, 20.0, 80.0, 20.0, 80.0, 40.0, 10.0, 40.0]


def test_extract_docx_pages_prefers_region_restore_items():
    payload = {
        "pages": [{
            "page_no": 1,
            "lines": [
                {"text": "global", "confidence": 0.99, "x1": 10, "y1": 20, "x2": 80, "y2": 40},
            ],
            "restore": {
                "region_restore": {
                    "regions": [{
                        "items": [
                            {"text": "restored", "confidence": 0.98, "rect": [100, 120, 180, 140]},
                        ],
                    }],
                },
            },
        }],
    }

    pages = extract_docx_pages(payload)

    assert pages[0]["ocr_lines"][0]["text"] == "restored"


def test_output_path_defaults_to_input_stem_with_docx_suffix():
    class Args:
        out = ""

    assert output_path(Args(), Path(r"D:\sample\page.png")) == Path(r"D:\sample\page.docx")


def test_build_best_docx_uses_pred_html_table_when_available(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "structure": {
                "raw_results": [{
                    "res": {
                        "layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 1000, 700]}]},
                        "overall_ocr_res": {
                            "rec_texts": ["outside", "inside"],
                            "rec_boxes": [[10, 10, 80, 30], [20, 120, 80, 140]],
                            "rec_scores": [0.99, 0.99],
                        },
                        "table_res_list": [{
                            "pred_html": "<table><tr><td>A</td><td colspan=\"2\">B</td></tr><tr><td colspan=\"3\"></td></tr><tr><td>C</td><td>D</td><td>E</td></tr></table>",
                        }],
                    },
                }],
            },
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "semantic.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2
    assert len(doc.tables[0].columns) == 3
    assert [paragraph.text for paragraph in doc.paragraphs] == ["outside"]


def test_build_best_docx_uses_plain_rows_for_non_table_documents(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1000,
            "height": 1400,
            "lines": [
                {"text": "title", "confidence": 0.99, "bbox": [100, 80, 240, 80, 240, 110, 100, 110]},
                {"text": "left", "confidence": 0.99, "bbox": [100, 180, 180, 180, 180, 210, 100, 210]},
                {"text": "right", "confidence": 0.99, "bbox": [700, 180, 800, 180, 800, 210, 700, 210]},
            ],
            "restore": {
                "region_restore": {
                    "regions": [
                        {"index": 1, "rect": [80, 70, 900, 260], "items": [
                            {"text": "region title", "rect": [100, 80, 260, 110]},
                        ]},
                    ],
                },
            },
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": []}}}]},
        }],
    }
    pages = extract_docx_pages(payload)
    output = tmp_path / "plain.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert len(doc.tables) == 0
    assert [paragraph.text for paragraph in doc.paragraphs] == ["title", "left  right"]


def test_plain_rows_preserve_vertical_ocr_boxes(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1000,
            "height": 1400,
            "lines": [
                {"text": "财税中国", "confidence": 0.99, "bbox": [80, 100, 110, 100, 110, 300, 80, 300]},
                {"text": "horizontal", "confidence": 0.99, "bbox": [180, 120, 360, 120, 360, 150, 180, 150]},
            ],
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": []}}}]},
        }],
    }
    pages = extract_docx_pages(payload)
    output = tmp_path / "plain_vertical.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert len(doc.tables) == 0
    assert doc.paragraphs[0].text == "财\n税\n中\n国  horizontal"


def test_plain_docx_can_insert_top_complex_image_block(tmp_path):
    from PIL import Image

    source = tmp_path / "source.png"
    Image.new("RGB", (1000, 1400), "white").save(source)
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1000,
            "height": 1400,
            "lines": [
                {"text": "竖向一", "confidence": 0.99, "bbox": [100, 100, 130, 100, 130, 360, 100, 360]},
                {"text": "竖向二", "confidence": 0.99, "bbox": [700, 100, 730, 100, 730, 360, 700, 360]},
                {"text": "body", "confidence": 0.99, "bbox": [120, 800, 220, 800, 220, 830, 120, 830]},
            ],
            "structure": {
                "raw_results": [{
                    "res": {
                        "layout_det_res": {"boxes": [
                            {"label": "image", "coordinate": [80, 80, 780, 500]},
                        ]},
                    },
                }],
            },
        }],
    }
    output = tmp_path / "image_block.docx"
    output.write_bytes(build_plain_docx_with_image_blocks(payload, source))

    doc = Document(str(output))

    assert len(doc.inline_shapes) == 1
    assert [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()] == ["body"]


def test_build_best_docx_prefers_coordinate_regions_and_keeps_footer_after_table(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [
                        {"index": 1, "rect": [0, 100, 100, 200], "items": [{"text": "buyer", "rect": [10, 120, 80, 140]}]},
                        {"index": 2, "rect": [100, 100, 300, 200], "items": [{"text": "seller", "rect": [110, 120, 180, 140]}]},
                        {"index": 3, "rect": [0, 200, 300, 300], "items": [{"text": "total", "rect": [10, 220, 80, 240]}]},
                    ],
                },
            },
            "structure": {
                "raw_results": [{
                    "res": {
                        "layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 300, 300]}]},
                        "overall_ocr_res": {
                            "rec_texts": ["footer"],
                            "rec_boxes": [[10, 340, 80, 360]],
                            "rec_scores": [0.99],
                        },
                        "table_res_list": [{
                            "pred_html": "<table><tr><td>pred-html</td></tr></table>",
                        }],
                    },
                }],
            },
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "regions.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert len(doc.tables) == 1
    assert doc.tables[0].cell(0, 0).text == "buyer"
    assert doc.tables[0].cell(0, 1).text == "seller"
    assert doc.tables[0].cell(1, 0).text == "total"
    assert [paragraph.text for paragraph in doc.paragraphs] == ["footer"]


def test_coordinate_region_cells_keep_item_rows(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [{
                        "index": 1,
                        "rect": [0, 100, 300, 300],
                        "items": [
                            {"text": "top-left", "rect": [10, 120, 80, 140]},
                            {"text": "top-right", "rect": [100, 120, 180, 140]},
                            {"text": "bottom", "rect": [10, 220, 80, 240]},
                        ],
                    }],
                },
            },
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 300, 300]}]}}}]},
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "region_rows.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert doc.tables[0].cell(0, 0).text == "top-lefttop-right\nbottom"


def test_dense_region_uses_nested_table_from_item_coordinates(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [{
                        "index": 1,
                        "rect": [0, 100, 600, 300],
                        "items": [
                            {"text": "H1", "rect": [0, 100, 100, 130]},
                            {"text": "H2", "rect": [200, 100, 240, 130]},
                            {"text": "H3", "rect": [300, 100, 340, 130]},
                            {"text": "H4", "rect": [400, 100, 440, 130]},
                            {"text": "H5", "rect": [500, 100, 580, 130]},
                            {"text": "R1C1", "rect": [0, 150, 120, 180]},
                            {"text": "R1C2", "rect": [200, 150, 240, 180]},
                            {"text": "R1C3", "rect": [300, 150, 340, 180]},
                            {"text": "R1C4", "rect": [400, 150, 440, 180]},
                            {"text": "R1C5", "rect": [500, 150, 540, 180]},
                        ],
                    }],
                },
            },
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 600, 300]}]}}}]},
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "detail_nested.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))
    nested = doc.tables[0].cell(0, 0).tables[0]

    assert len(nested.rows) == 2
    assert len(nested.columns) == 5
    assert nested.cell(0, 0).text == "H1"
    assert nested.cell(1, 0).text == "R1C1"
    assert nested.cell(1, 3).text == "R1C4"
    assert "w:val=\"nil\"" in nested._tbl.xml


def test_single_row_dense_region_uses_borderless_nested_positioning_table(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [{
                        "index": 1,
                        "rect": [0, 100, 700, 180],
                        "items": [
                            {"text": "C1", "rect": [0, 100, 60, 130]},
                            {"text": "C2", "rect": [100, 100, 220, 130]},
                            {"text": "C3", "rect": [260, 100, 340, 130]},
                            {"text": "C4", "rect": [380, 100, 440, 130]},
                            {"text": "C5", "rect": [480, 100, 540, 130]},
                            {"text": "C6", "rect": [580, 100, 620, 130]},
                            {"text": "C7", "rect": [640, 100, 700, 130]},
                        ],
                    }],
                },
            },
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 700, 180]}]}}}]},
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "travel_nested.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))
    nested = doc.tables[0].cell(0, 0).tables[0]

    assert len(nested.columns) == 7
    assert nested.cell(0, 0).text == "C1"
    assert nested.cell(0, 6).text == "C7"
    assert "w:val=\"nil\"" in nested._tbl.xml


def test_narrow_region_label_is_written_vertically(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [{
                        "index": 1,
                        "rect": [0, 100, 40, 300],
                        "items": [{"text": "纵向标签", "rect": [5, 120, 35, 280]}],
                    }],
                },
            },
            "structure": {"raw_results": [{"res": {"layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 40, 300]}]}}}]},
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "vertical.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert doc.tables[0].cell(0, 0).text == "纵\n向\n标\n签"


def test_outside_text_uses_paragraphs_to_avoid_helper_gridlines(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {"region_restore": {"regions": [{"index": 1, "rect": [0, 100, 1000, 300], "items": []}]}},
            "structure": {
                "raw_results": [{
                    "res": {
                        "layout_det_res": {"boxes": [{"label": "table", "coordinate": [0, 100, 1000, 300]}]},
                        "overall_ocr_res": {
                            "rec_texts": ["left", "right"],
                            "rec_boxes": [[10, 20, 80, 40], [900, 20, 980, 40]],
                            "rec_scores": [0.99, 0.99],
                        },
                    },
                }],
            },
        }],
    }
    pages = [{"page": 1, "ocr_lines": [{"text": "grid", "confidence": 0.99, "bbox": [1, 1, 2, 1, 2, 2, 1, 2]}]}]
    output = tmp_path / "outside.docx"
    output.write_bytes(build_best_docx(payload, pages, grid_cols=36, grid_rows=48))

    doc = Document(str(output))

    assert len(doc.tables) == 0
    assert [paragraph.text for paragraph in doc.paragraphs] == ["left  right"]


def test_build_table_only_docx_prefers_large_regions_over_cell_boxes(tmp_path):
    payload = {
        "pages": [{
            "page_no": 1,
            "width": 1200,
            "height": 800,
            "restore": {
                "region_restore": {
                    "regions": [
                        {"index": 1, "rect": [0, 100, 100, 200], "items": []},
                        {"index": 2, "rect": [100, 100, 300, 200], "items": []},
                        {"index": 3, "rect": [0, 200, 300, 300], "items": []},
                    ],
                },
            },
            "structure": {
                "summary": {
                    "tables": [{
                        "cell_boxes": [
                            [0, 0, 10, 10],
                            [10, 0, 20, 10],
                            [20, 0, 30, 10],
                            [30, 0, 40, 10],
                        ],
                    }],
                },
            },
        }],
    }
    output = tmp_path / "table_only.docx"
    output.write_bytes(build_table_only_docx(payload))

    doc = Document(str(output))

    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2
    assert len(doc.tables[0].columns) == 2
    assert all(cell.text == "" for row in doc.tables[0].rows for cell in row.cells)
