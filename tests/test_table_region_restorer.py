import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from table_region_restorer import build_region_docx, build_region_paint_html, restore_table_regions


def test_restore_table_regions_assigns_text_to_detected_region():
    regions = [{"label": "cell", "score": 0.9, "coordinate": [0, 260, 200, 360]}]
    table_ocr = {
        "rec_texts": ["hello"],
        "rec_scores": [0.99],
        "rec_boxes": [[20, 280, 80, 310]],
    }

    restored = restore_table_regions(regions, table_ocr)

    assert restored["regions"][0]["text"] == "hello"


def test_restore_table_regions_uses_supplemental_longer_ocr():
    regions = [{"label": "cell", "score": 0.9, "coordinate": [0, 260, 900, 460]}]
    table_ocr = {
        "rec_texts": ["人", "计", "EM"],
        "rec_scores": [0.9, 0.9, 0.9],
        "rec_boxes": [[100, 380, 130, 410], [180, 380, 210, 410], [220, 380, 800, 410]],
    }
    supplemental = {
        "rec_texts": ["统一社会信用代码/纳税人识别号：91441900762909683X"],
        "rec_scores": [0.99],
        "rec_boxes": [[90, 378, 820, 412]],
    }

    restored = restore_table_regions(regions, table_ocr, supplemental_ocr=supplemental)

    assert "统一社会信用代码" in restored["regions"][0]["text"]
    assert "计" not in restored["regions"][0]["text"]
    assert "EM" not in restored["regions"][0]["text"]


def test_build_region_docx_uses_detail_nested_table(tmp_path):
    regions = [{
        "index": 1,
        "kind": "detail",
        "rect": [0, 0, 600, 200],
        "text": "",
        "items": [
            {"text": "项目名称", "rect": [0, 0, 80, 20]},
            {"text": "单价", "rect": [100, 0, 160, 20]},
            {"text": "数量", "rect": [200, 0, 260, 20]},
            {"text": "金额", "rect": [300, 0, 360, 20]},
            {"text": "税率/征收率", "rect": [400, 0, 500, 20]},
            {"text": "税额", "rect": [520, 0, 580, 20]},
            {"text": "17.09", "rect": [100, 40, 160, 60]},
        ],
    }]
    path = tmp_path / "regions.docx"
    path.write_bytes(build_region_docx({"regions": regions}))

    doc = Document(str(path))
    texts = _all_cell_texts(doc.tables)

    assert "项目名称" in texts
    assert "17.09" in texts


def test_build_region_paint_html_draws_without_background_image():
    restored = {
        "regions": [{
            "index": 1,
            "kind": "region",
            "rect": [0, 260, 200, 360],
            "text": "cell text",
            "items": [{"text": "cell text", "rect": [20, 280, 80, 310]}],
        }],
        "ocr_items": [],
        "page_blocks": [{"label": "text", "text": "invoice no", "rect": [300, 20, 500, 60]}],
    }

    output = build_region_paint_html(restored)

    assert "page-image" not in output
    assert "paint-region" in output
    assert "cell text" in output
    assert "invoice no" in output


def test_build_region_paint_html_uses_assigned_region_items_once():
    restored = {
        "regions": [{
            "index": 1,
            "kind": "region",
            "rect": [0, 260, 200, 360],
            "text": "canonical text",
            "items": [{"text": "canonical text", "rect": [20, 280, 120, 310]}],
        }],
        "ocr_items": [
            {"text": "canonical text", "rect": [20, 280, 120, 310]},
            {"text": "stale duplicate", "rect": [22, 282, 122, 312]},
        ],
        "page_blocks": [],
    }

    output = build_region_paint_html(restored)

    assert output.count("canonical text") == 1
    assert "stale duplicate" not in output


def test_build_region_paint_html_marks_tall_cjk_as_vertical():
    restored = {
        "regions": [],
        "ocr_items": [{"text": "备注", "rect": [10, 10, 50, 110]}],
        "page_blocks": [],
    }

    output = build_region_paint_html(restored)

    assert "paint-text vertical" in output
    assert "writing-mode: vertical-rl" in output


def _all_cell_texts(tables):
    texts = []
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
                texts.extend(_all_cell_texts(cell.tables))
    return texts
