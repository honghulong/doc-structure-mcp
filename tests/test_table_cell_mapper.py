import json
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from table_cell_mapper import build_cell_docx, build_cell_overlay_html, load_structure_result, map_table_cells


def test_load_structure_result_unwraps_paddleocr_res(tmp_path):
    path = tmp_path / "structure.json"
    path.write_text(json.dumps({"res": {"table_res_list": []}}), encoding="utf-8")

    result = load_structure_result(path)

    assert result == {"table_res_list": []}


def test_map_table_cells_assigns_ocr_text_by_cell_box():
    result = {
        "table_res_list": [{
            "cell_box_list": [
                [0, 0, 100, 50],
                [100, 0, 200, 50],
            ],
            "table_ocr_pred": {
                "rec_texts": ["left", "right"],
                "rec_scores": [0.99, 0.98],
                "rec_boxes": [
                    [10, 10, 80, 40],
                    [120, 10, 180, 40],
                ],
            },
        }]
    }

    mapping = map_table_cells(result)
    cells = mapping["tables"][0]["cells"]

    assert mapping["tables"][0]["cell_count"] == 2
    assert cells[0]["text"] == "left"
    assert cells[1]["text"] == "right"


def test_map_table_cells_keeps_non_table_text_blocks():
    result = {
        "parsing_res_list": [
            {"block_label": "text", "block_content": "invoice no", "block_bbox": [10, 10, 90, 30]},
            {"block_label": "table", "block_content": "", "block_bbox": [0, 50, 200, 120]},
        ],
        "table_res_list": [{"cell_box_list": [], "table_ocr_pred": {}}],
    }

    mapping = map_table_cells(result)

    assert mapping["blocks"][0]["text"] == "invoice no"


def test_cell_overlay_html_contains_cell_ids_and_text():
    mapping = {
        "tables": [{
            "cells": [{"index": 1, "rect": [0, 0, 100, 50], "text": "hello", "items": []}]
        }]
    }

    output = build_cell_overlay_html(mapping)

    assert "cell-id" in output
    assert "hello" in output


def test_cell_overlay_html_contains_non_table_blocks():
    mapping = {
        "blocks": [{"label": "text", "rect": [0, 0, 100, 30], "text": "invoice no"}],
        "tables": [{"cells": []}],
    }

    output = build_cell_overlay_html(mapping)

    assert "block-label" in output
    assert "invoice no" in output


def test_build_cell_docx_creates_editable_table(tmp_path):
    mapping = {
        "tables": [{
            "cells": [
                {"index": 1, "rect": [0, 0, 100, 50], "text": "left", "items": []},
                {"index": 2, "rect": [100, 0, 200, 50], "text": "right", "items": []},
            ]
        }]
    }
    path = tmp_path / "cells.docx"
    path.write_bytes(build_cell_docx(mapping))

    doc = Document(str(path))
    texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]

    assert "left" in texts
    assert "right" in texts


def test_build_cell_docx_preserves_sparse_x_columns(tmp_path):
    mapping = {
        "tables": [{
            "cells": [
                {"index": 1, "rect": [0, 0, 50, 30], "text": "left", "items": []},
                {"index": 2, "rect": [200, 0, 250, 30], "text": "right", "items": []},
                {"index": 3, "rect": [200, 40, 250, 70], "text": "below-right", "items": []},
            ]
        }]
    }
    path = tmp_path / "sparse.docx"
    path.write_bytes(build_cell_docx(mapping))

    doc = Document(str(path))
    table = doc.tables[0]

    assert len(table.columns) >= 3
    assert table.cell(0, 0).text == "left"
    assert table.cell(0, 2).text == "right"
    assert table.cell(1, 2).text == "below-right"
