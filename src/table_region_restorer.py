import base64
import html
import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt


DETAIL_HEADERS = ["项目名称", "单价", "数量", "金额", "税率/征收率", "税额"]


def load_table_ocr(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("res"), dict):
        data = data["res"]
    table = data.get("table_res_list", [{}])[0]
    return table.get("table_ocr_pred", {})


def load_overall_ocr(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("res"), dict):
        data = data["res"]
    return data.get("overall_ocr_res", {})


def load_page_blocks(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("res"), dict):
        data = data["res"]
    blocks = []
    for block in data.get("parsing_res_list", []):
        label = block.get("block_label", "")
        if label in {"table", "seal", "image"}:
            continue
        text = str(block.get("block_content", "")).strip()
        bbox = block.get("block_bbox")
        if text and bbox:
            blocks.append({"label": label, "text": text, "rect": _rect(bbox)})
    return blocks


def load_detected_regions(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("res"), dict):
        data = data["res"]
    return data.get("boxes", [])


def restore_table_regions(regions, table_ocr, supplemental_ocr=None):
    detected = [
        region for index, box in enumerate(regions, start=1)
        if _is_table_body_region(region := _normalize_region(index, box))
    ]
    ocr_items = _merge_ocr_items(_ocr_items(table_ocr), _ocr_items(supplemental_ocr or {}))
    for item in ocr_items:
        region = _best_region(item["rect"], detected)
        if region:
            region["items"].append(item)
    for region in detected:
        region["items"].sort(key=lambda item: (item["rect"][1], item["rect"][0]))
        region["text"] = _region_text(region)
        region["kind"] = _region_kind(region)
    detected.sort(key=lambda item: (item["rect"][1], item["rect"][0]))
    return {"regions": detected, "ocr_items": ocr_items}


def build_region_overlay_html(restored, image_path=None):
    src, width, height = _image_info(image_path)
    if not width or not height:
        width, height = _infer_size(restored.get("regions", []))
    parts = []
    for region in restored.get("regions", []):
        x1, y1, x2, y2 = region["rect"]
        style = (
            f"left:{100*x1/width:.4f}%;top:{100*y1/height:.4f}%;"
            f"width:{100*(x2-x1)/width:.4f}%;height:{100*(y2-y1)/height:.4f}%;"
        )
        text = html.escape(region.get("text", ""))
        parts.append(
            f'<div class="region {region["kind"]}" style="{style}" title="{text}">'
            f'<span class="region-id">R{region["index"]} {html.escape(region["kind"])}</span>'
            f'<span class="region-text">{text}</span></div>'
        )
    background = f'<img class="page-image" src="{src}" alt="page">' if src else ""
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>table region restore</title>
<style>
body {{ margin: 0; padding: 24px; background: #f3f4f6; font-family: Arial, "Microsoft YaHei", sans-serif; }}
.page {{ position: relative; max-width: 1180px; margin: 0 auto; background: white; aspect-ratio: {width}/{height}; }}
.page-image {{ position: absolute; inset: 0; width: 100%; height: 100%; object-fit: contain; }}
.region {{ position: absolute; box-sizing: border-box; border: 3px solid rgba(220,38,38,.78); background: rgba(254,226,226,.18); overflow: hidden; }}
.region.detail {{ border-color: rgba(37,99,235,.84); background: rgba(219,234,254,.20); }}
.region-id {{ position: absolute; left: 3px; top: 2px; color: #991b1b; font-size: 11px; font-weight: 700; }}
.region.detail .region-id {{ color: #1d4ed8; }}
.region-text {{ display: block; padding: 17px 5px 3px; color: #111827; font-size: 12px; line-height: 1.18; white-space: pre-wrap; }}
</style>
</head>
<body><section class="page">{background}{''.join(parts)}</section></body>
</html>
"""


def build_region_paint_html(restored, image_path=None):
    _, width, height = _image_info(image_path)
    if not width or not height:
        width, height = _infer_size(restored.get("regions", []))

    parts = []
    for region in restored.get("regions", []):
        x1, y1, x2, y2 = region["rect"]
        parts.append(
            f'<div class="paint-region {region["kind"]}" style="{_css_rect(x1, y1, x2, y2, width, height)}">'
            f'<span class="paint-region-label">R{region["index"]}</span></div>'
        )
    for item in _paint_region_items(restored):
        x1, y1, x2, y2 = item["rect"]
        text = html.escape(item.get("text", ""))
        text_class = " vertical" if _looks_vertical_text(item) else ""
        parts.append(
            f'<div class="paint-text{text_class}" style="{_css_rect(x1, y1, x2, y2, width, height)}">{text}</div>'
        )
    for block in restored.get("page_blocks", []):
        x1, y1, x2, y2 = block["rect"]
        text = html.escape(block.get("text", ""))
        label = html.escape(block.get("label", "block"))
        parts.append(
            f'<div class="paint-block" title="{label}" style="{_css_rect(x1, y1, x2, y2, width, height)}">{text}</div>'
        )

    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>painted table restore</title>
<style>
body {{ margin: 0; padding: 24px; background: #e5e7eb; font-family: Arial, "Microsoft YaHei", sans-serif; }}
.paint-page {{ position: relative; max-width: 1180px; margin: 0 auto; background: white; aspect-ratio: {width}/{height}; box-shadow: 0 2px 14px rgba(15,23,42,.18); overflow: hidden; }}
.paint-region {{ position: absolute; box-sizing: border-box; border: 2px solid #991b1b; background: transparent; }}
.paint-region.detail {{ border-color: #1d4ed8; }}
.paint-region.remark {{ border-color: #be123c; }}
.paint-region-label {{ position: absolute; left: 2px; top: 1px; color: #991b1b; font-size: 10px; font-weight: 700; }}
.paint-region.detail .paint-region-label {{ color: #1d4ed8; }}
.paint-text {{ position: absolute; box-sizing: border-box; color: #111827; font-size: 14px; line-height: 1.05; white-space: nowrap; overflow: visible; }}
.paint-text.vertical {{ writing-mode: vertical-rl; text-orientation: upright; display: flex; align-items: center; justify-content: center; white-space: normal; line-height: 1.08; }}
.paint-block {{ position: absolute; box-sizing: border-box; color: #7f1d1d; font-size: 15px; line-height: 1.1; white-space: pre-wrap; overflow: visible; }}
</style>
</head>
<body><section class="paint-page">{''.join(parts)}</section></body>
</html>
"""


def _paint_region_items(restored):
    regions = restored.get("regions", [])
    if not regions:
        return restored.get("ocr_items", [])
    items = []
    for region in regions:
        items.extend(region.get("items", []))
    return items


def build_region_docx(restored):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(9)

    rows = _cluster_region_rows(restored.get("regions", []))
    max_cols = max((len(row) for row in rows), default=1)
    table = doc.add_table(rows=max(1, len(rows)), cols=max_cols)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index, region in enumerate(row):
            cell = table.cell(row_index, col_index)
            if region["kind"] == "detail":
                _fill_detail_region(cell, region)
            else:
                cell.text = region.get("text", "")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "SimSun"
                    run.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _fill_detail_region(word_cell, region):
    word_cell.text = ""
    items = region.get("items", [])
    header_items = [item for item in items if item["text"] in DETAIL_HEADERS]
    col_centers = _detail_col_centers(header_items)
    if len(col_centers) < len(DETAIL_HEADERS):
        word_cell.text = region.get("text", "")
        return
    rows = _cluster_ocr_rows(items)
    nested = word_cell.add_table(rows=max(1, len(rows)), cols=len(DETAIL_HEADERS))
    nested.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for item in row:
            col_index = _nearest_index((item["rect"][0] + item["rect"][2]) / 2, col_centers)
            target = nested.cell(row_index, col_index)
            if target.text:
                target.text += " "
            target.text += item["text"]


def _normalize_region(index, box):
    return {
        "index": index,
        "score": float(box.get("score", 0)),
        "label": box.get("label", ""),
        "rect": [float(v) for v in box.get("coordinate", [])],
        "items": [],
        "text": "",
        "kind": "region",
    }


def _is_table_body_region(region):
    x1, y1, x2, y2 = region["rect"]
    return y1 >= 240 and x2 > 30 and y2 <= 1130


def _ocr_items(pred):
    texts = pred.get("rec_texts", []) if isinstance(pred, dict) else []
    boxes = pred.get("rec_boxes") or pred.get("rec_polys") or []
    scores = pred.get("rec_scores", [])
    items = []
    for index, text in enumerate(texts):
        text = str(text).strip()
        if not text or index >= len(boxes):
            continue
        rect = _rect(boxes[index])
        items.append({
            "text": text,
            "rect": rect,
            "confidence": float(scores[index]) if index < len(scores) else 0.0,
        })
    return items


def _merge_ocr_items(primary, supplemental):
    merged = list(primary)
    for item in supplemental:
        if _is_duplicate_item(item, merged):
            continue
        merged = [
            existing for existing in merged
            if not _contains_rect(item["rect"], existing["rect"]) or len(item["text"]) <= len(existing["text"])
        ]
        merged.append(item)
    return merged


def _is_duplicate_item(item, existing_items):
    for existing in existing_items:
        if item["text"] == existing["text"] and _overlap_ratio(item["rect"], existing["rect"]) > 0.8:
            return True
    return False


def _contains_rect(outer, inner):
    return outer[0] <= inner[0] and outer[1] <= inner[1] and outer[2] >= inner[2] and outer[3] >= inner[3]


def _best_region(rect, regions):
    best = None
    best_score = 0.0
    for region in regions:
        score = _overlap_ratio(rect, region["rect"])
        if score > best_score:
            best = region
            best_score = score
    if best and (best_score >= 0.35 or _center_inside(rect, best["rect"])):
        return best
    return None


def _region_text(region):
    rows = _cluster_ocr_rows(region.get("items", []))
    return "\n".join("  ".join(item["text"] for item in row) for row in rows if row)


def _region_kind(region):
    texts = {item["text"] for item in region.get("items", [])}
    if any(text in texts for text in DETAIL_HEADERS):
        return "detail"
    if "备注" in texts:
        return "remark"
    return "region"


def _looks_vertical_text(item):
    text = item.get("text", "").strip()
    x1, y1, x2, y2 = item["rect"]
    width = max(1.0, x2 - x1)
    height = max(1.0, y2 - y1)
    if len(text) < 2 or len(text) > 8:
        return False
    if height / width < 1.8:
        return False
    return _cjk_ratio(text) >= 0.75


def _cjk_ratio(text):
    chars = [char for char in text if not char.isspace()]
    if not chars:
        return 0.0
    cjk_count = sum(1 for char in chars if "\u4e00" <= char <= "\u9fff")
    return cjk_count / len(chars)


def _cluster_region_rows(regions):
    rows = []
    for region in sorted(regions, key=lambda item: ((item["rect"][1] + item["rect"][3]) / 2, item["rect"][0])):
        cy = (region["rect"][1] + region["rect"][3]) / 2
        height = max(1.0, region["rect"][3] - region["rect"][1])
        threshold = max(22.0, height * 0.45)
        if rows and abs(cy - _row_center(rows[-1])) <= threshold:
            rows[-1].append(region)
        else:
            rows.append([region])
    for row in rows:
        row.sort(key=lambda item: item["rect"][0])
    return rows


def _cluster_ocr_rows(items):
    rows = []
    for item in sorted(items, key=lambda value: ((value["rect"][1] + value["rect"][3]) / 2, value["rect"][0])):
        cy = (item["rect"][1] + item["rect"][3]) / 2
        height = max(1.0, item["rect"][3] - item["rect"][1])
        threshold = max(9.0, height * 0.6)
        if rows and abs(cy - _ocr_row_center(rows[-1])) <= threshold:
            rows[-1].append(item)
        else:
            rows.append([item])
    for row in rows:
        row.sort(key=lambda value: value["rect"][0])
    return rows


def _detail_col_centers(header_items):
    by_name = {item["text"]: (item["rect"][0] + item["rect"][2]) / 2 for item in header_items}
    return [by_name[name] for name in DETAIL_HEADERS if name in by_name]


def _nearest_index(value, centers):
    return min(range(len(centers)), key=lambda index: abs(value - centers[index]))


def _row_center(row):
    return sum((item["rect"][1] + item["rect"][3]) / 2 for item in row) / len(row)


def _ocr_row_center(row):
    return sum((item["rect"][1] + item["rect"][3]) / 2 for item in row) / len(row)


def _rect(box):
    if len(box) == 4 and all(isinstance(v, (int, float)) for v in box):
        return [float(v) for v in box]
    points = [[float(p[0]), float(p[1])] for p in box if len(p) >= 2]
    xs = [point[0] for point in points]
    ys = [point[1] for point in points]
    return [min(xs), min(ys), max(xs), max(ys)]


def _overlap_ratio(a, b):
    x1 = max(a[0], b[0])
    y1 = max(a[1], b[1])
    x2 = min(a[2], b[2])
    y2 = min(a[3], b[3])
    if x2 <= x1 or y2 <= y1:
        return 0.0
    return ((x2 - x1) * (y2 - y1)) / max(1.0, (a[2] - a[0]) * (a[3] - a[1]))


def _center_inside(a, b):
    cx = (a[0] + a[2]) / 2
    cy = (a[1] + a[3]) / 2
    return b[0] <= cx <= b[2] and b[1] <= cy <= b[3]


def _infer_size(regions):
    max_x = max((region["rect"][2] for region in regions), default=1)
    max_y = max((region["rect"][3] for region in regions), default=1)
    return int(max_x), int(max_y)


def _image_info(image_path):
    if not image_path:
        return "", None, None
    path = Path(image_path)
    if not path.exists():
        return "", None, None
    try:
        from PIL import Image
        with Image.open(path) as image:
            width, height = image.size
    except Exception:
        width = height = None
    data = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{data}", width, height


def _css_rect(x1, y1, x2, y2, width, height):
    return (
        f"left:{100*x1/width:.4f}%;top:{100*y1/height:.4f}%;"
        f"width:{100*(x2-x1)/width:.4f}%;height:{100*(y2-y1)/height:.4f}%;"
    )
