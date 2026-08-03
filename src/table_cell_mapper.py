import base64
import html
import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt


def load_structure_result(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, dict) and isinstance(data.get("res"), dict):
        return data["res"]
    return data


def map_table_cells(structure_result, min_overlap=0.45):
    """Map table OCR text boxes into PPStructure cell boxes."""
    table = _first_table(structure_result)
    if not table:
        return {"blocks": _non_table_blocks(structure_result), "tables": []}

    cells = [_normalize_cell(index, box) for index, box in enumerate(table.get("cell_box_list", []), start=1)]
    ocr_items = _table_ocr_items(table.get("table_ocr_pred", {}))

    for item in ocr_items:
        best_cell = None
        best_score = 0.0
        item_rect = item["rect"]
        for cell in cells:
            score = _overlap_ratio(item_rect, cell["rect"])
            if score > best_score:
                best_score = score
                best_cell = cell
        if best_cell and (best_score >= min_overlap or _center_inside(item_rect, best_cell["rect"])):
            best_cell["items"].append({**item, "overlap": round(best_score, 4)})

    for cell in cells:
        cell["text"] = " ".join(item["text"] for item in sorted(cell["items"], key=lambda x: (x["rect"][1], x["rect"][0])) if item["text"])

    return {
        "blocks": _non_table_blocks(structure_result),
        "tables": [{
            "cell_count": len(cells),
            "ocr_count": len(ocr_items),
            "cells": cells,
            "pred_html": table.get("pred_html", ""),
        }]
    }


def build_cell_overlay_html(mapping, image_path=None):
    image_src, width, height = _image_info(image_path)
    table = mapping["tables"][0] if mapping.get("tables") else {"cells": []}
    cells = table.get("cells", [])
    blocks = mapping.get("blocks", [])
    if not width or not height:
        width, height = _infer_size(cells + blocks)

    parts = []
    for block in blocks:
        x1, y1, x2, y2 = block["rect"]
        style = (
            f"left:{100*x1/width:.4f}%;top:{100*y1/height:.4f}%;"
            f"width:{100*(x2-x1)/width:.4f}%;height:{100*(y2-y1)/height:.4f}%;"
        )
        text = html.escape(block.get("text", ""))
        label = html.escape(block.get("label", "block"))
        parts.append(
            f'<div class="block" title="{label}" style="{style}">'
            f'<span class="block-label">{label}</span><span class="block-text">{text}</span></div>'
        )
    for cell in cells:
        x1, y1, x2, y2 = cell["rect"]
        style = (
            f"left:{100*x1/width:.4f}%;top:{100*y1/height:.4f}%;"
            f"width:{100*(x2-x1)/width:.4f}%;height:{100*(y2-y1)/height:.4f}%;"
        )
        text = html.escape(cell.get("text", ""))
        title = html.escape(f"cell {cell['index']} | {text}")
        parts.append(
            f'<div class="cell" title="{title}" style="{style}">'
            f'<span class="cell-id">{cell["index"]}</span><span class="cell-text">{text}</span></div>'
        )

    background = f'<img class="page-image" src="{image_src}" alt="page">' if image_src else ""
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>table cell overlay</title>
<style>
body {{ margin: 0; padding: 24px; background: #f3f4f6; font-family: Arial, "Microsoft YaHei", sans-serif; }}
.page {{ position: relative; max-width: 1180px; margin: 0 auto; background: white; aspect-ratio: {width}/{height}; }}
.page-image {{ position: absolute; inset: 0; width: 100%; height: 100%; object-fit: contain; }}
.cell {{ position: absolute; box-sizing: border-box; border: 2px solid rgba(37, 99, 235, .72); background: rgba(219, 234, 254, .22); overflow: hidden; }}
.cell-id {{ position: absolute; left: 2px; top: 1px; color: #1d4ed8; font-size: 11px; font-weight: 700; }}
.cell-text {{ display: block; padding: 14px 3px 2px; color: #111827; font-size: 12px; line-height: 1.15; }}
.block {{ position: absolute; box-sizing: border-box; border: 2px solid rgba(5, 150, 105, .76); background: rgba(209, 250, 229, .28); overflow: hidden; }}
.block-label {{ position: absolute; left: 2px; top: 1px; color: #047857; font-size: 10px; font-weight: 700; }}
.block-text {{ display: block; padding: 13px 3px 2px; color: #064e3b; font-size: 12px; line-height: 1.15; }}
</style>
</head>
<body>
<section class="page">{background}{''.join(parts)}</section>
</body>
</html>
"""


def build_cell_docx(mapping):
    """Build an editable DOCX table from mapped cell text."""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "SimSun"
    style.font.size = Pt(9)

    for table_index, mapped_table in enumerate(mapping.get("tables", []), start=1):
        if table_index > 1:
            doc.add_page_break()
        for block in sorted(mapping.get("blocks", []), key=lambda item: (item["rect"][1], item["rect"][0])):
            if block.get("text"):
                doc.add_paragraph(block["text"])
        rows = _cluster_cell_rows(mapped_table.get("cells", []))
        if not rows:
            doc.add_paragraph("")
            continue
        col_edges = _cluster_axis_edges(mapped_table.get("cells", []), axis="x")
        max_cols = max(1, len(col_edges) - 1)
        table = doc.add_table(rows=len(rows), cols=max_cols)
        table.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for cell_data in row:
                col_index = _edge_index((cell_data["rect"][0] + cell_data["rect"][2]) / 2, col_edges)
                word_cell = table.cell(row_index, col_index)
                word_cell.text = cell_data.get("text", "")
                for paragraph in word_cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "SimSun"
                        run.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _first_table(result):
    tables = result.get("table_res_list") if isinstance(result, dict) else None
    if isinstance(tables, list) and tables:
        return tables[0]
    return None


def _non_table_blocks(result):
    blocks = []
    parsing = result.get("parsing_res_list") if isinstance(result, dict) else None
    if not isinstance(parsing, list):
        return blocks
    for block in parsing:
        label = block.get("block_label", "")
        if label in {"table", "seal", "image"}:
            continue
        text = str(block.get("block_content", "")).strip()
        if not text:
            continue
        bbox = block.get("block_bbox")
        if not bbox:
            continue
        blocks.append({
            "label": label,
            "text": text,
            "rect": _rect(bbox),
        })
    return blocks


def _normalize_cell(index, box):
    rect = _rect(box)
    return {"index": index, "rect": rect, "items": [], "text": ""}


def _table_ocr_items(pred):
    texts = pred.get("rec_texts", []) if isinstance(pred, dict) else []
    boxes = pred.get("rec_boxes") or pred.get("rec_polys") or []
    scores = pred.get("rec_scores", [])
    items = []
    for index, text in enumerate(texts):
        if index >= len(boxes):
            continue
        rect = _rect(boxes[index])
        score = float(scores[index]) if index < len(scores) else 0.0
        items.append({"text": str(text).strip(), "confidence": score, "rect": rect})
    return items


def _rect(box):
    if len(box) == 4 and all(isinstance(v, (int, float)) for v in box):
        x1, y1, x2, y2 = [float(v) for v in box]
        return [x1, y1, x2, y2]
    points = [[float(p[0]), float(p[1])] for p in box if len(p) >= 2]
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    return [min(xs), min(ys), max(xs), max(ys)]


def _overlap_ratio(a, b):
    x1 = max(a[0], b[0])
    y1 = max(a[1], b[1])
    x2 = min(a[2], b[2])
    y2 = min(a[3], b[3])
    if x2 <= x1 or y2 <= y1:
        return 0.0
    area = (x2 - x1) * (y2 - y1)
    item_area = max(1.0, (a[2] - a[0]) * (a[3] - a[1]))
    return area / item_area


def _center_inside(a, b):
    cx = (a[0] + a[2]) / 2
    cy = (a[1] + a[3]) / 2
    return b[0] <= cx <= b[2] and b[1] <= cy <= b[3]


def _infer_size(cells):
    max_x = max((cell["rect"][2] for cell in cells), default=1)
    max_y = max((cell["rect"][3] for cell in cells), default=1)
    return int(max_x), int(max_y)


def _image_info(image_path):
    if not image_path:
        return "", None, None
    path = Path(image_path)
    if not path.exists():
        return "", None, None
    try:
        from PIL import Image
        with Image.open(path) as image:
            width, height = image.size
    except Exception:
        width = height = None
    mime = "image/jpeg" if path.suffix.lower() in {".jpg", ".jpeg"} else "image/png"
    data = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{data}", width, height


def _cluster_cell_rows(cells):
    rows = []
    for cell in sorted(cells, key=lambda item: ((item["rect"][1] + item["rect"][3]) / 2, item["rect"][0])):
        cy = (cell["rect"][1] + cell["rect"][3]) / 2
        height = max(1.0, cell["rect"][3] - cell["rect"][1])
        threshold = max(10.0, height * 0.7)
        if rows and abs(cy - _cell_row_center(rows[-1])) <= threshold:
            rows[-1].append(cell)
        else:
            rows.append([cell])
    for row in rows:
        row.sort(key=lambda item: item["rect"][0])
    return rows


def _cell_row_center(row):
    return sum((cell["rect"][1] + cell["rect"][3]) / 2 for cell in row) / len(row)


def _cluster_axis_edges(cells, axis):
    indexes = (0, 2) if axis == "x" else (1, 3)
    values = []
    for cell in cells:
        values.extend([cell["rect"][indexes[0]], cell["rect"][indexes[1]]])
    clusters = []
    for value in sorted(values):
        if clusters and abs(value - clusters[-1][-1]) <= 18:
            clusters[-1].append(value)
        else:
            clusters.append([value])
    return [sum(cluster) / len(cluster) for cluster in clusters]


def _edge_index(value, edges):
    if len(edges) < 2:
        return 0
    centers = [(edges[i] + edges[i + 1]) / 2 for i in range(len(edges) - 1)]
    return min(range(len(centers)), key=lambda index: abs(value - centers[index]))
